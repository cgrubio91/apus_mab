"""Carga el directorio de proveedores del IDU y su Banco de Precios de Referencia (BPR).

Dos archivos del IDU que comparten la llave "Grupo de base de datos":

  * "Visor_directorio_de_proveedores_de_cotizaciones_*.xlsx"
      → tablas `proveedor` y `proveedor_grupo` (a quién pedirle cotización).
  * "*Visor_BPR_*.xlsx" (hoja "Inusmos", así, con el typo del original)
      → tabla `insumo_referencia_idu` (precio oficial por insumo).

Es idempotente: se puede volver a correr con una publicación nueva y actualiza
los registros existentes en vez de duplicarlos.

Uso:
    # Ver qué haría, sin escribir nada:
    python scripts/cargar_proveedores_idu.py --directorio "ruta/Visor_directorio...xlsx" --dry-run

    # Cargar ambos:
    python scripts/cargar_proveedores_idu.py \
        --directorio "ruta/Visor_directorio...xlsx" \
        --bpr "ruta/1P. Visor_BPR...xlsx" \
        --periodo "2026-I Fase I"

Conexión: DB_HOST, DB_PORT, DB_USER, DB_PASSWORD, DB_NAME (o el .env del proyecto).
"""

import argparse
import logging
import sys
from pathlib import Path
from typing import Optional

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.infrastructure.database.connection import execute_query  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("proveedores_idu")

# Hoja y columnas del directorio (índices 0-based sobre la fila completa).
HOJA_DIRECTORIO = "Directorio General"
FILA_INICIO_DIRECTORIO = 5
COL = {
    "grupo": 2, "nombre": 3, "municipio": 4, "departamento": 5,
    "direccion": 6, "telefono": 7, "web_correo": 8, "contacto": 9, "cotizo": 15,
}

# Hoja del BPR (el archivo original trae "Inusmos" mal escrito).
HOJAS_BPR = ("Inusmos", "Insumos")
FILA_INICIO_BPR = 12
COL_BPR = {"origen": 1, "grupo": 2, "codigo": 3, "nombre": 4, "unidad": 5, "precio": 6}


def _texto(valor) -> Optional[str]:
    if valor is None:
        return None
    t = " ".join(str(valor).split())
    return t or None


def _numero(valor) -> Optional[float]:
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    limpio = str(valor).strip().replace("$", "").replace(" ", "").replace(",", "")
    try:
        return float(limpio)
    except ValueError:
        return None


def leer_directorio(ruta: Path) -> list[dict]:
    """Filas proveedor-grupo del directorio."""
    wb = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    if HOJA_DIRECTORIO not in wb.sheetnames:
        wb.close()
        raise ValueError(f"El archivo no tiene la hoja '{HOJA_DIRECTORIO}'")

    filas = []
    for fila in wb[HOJA_DIRECTORIO].iter_rows(min_row=FILA_INICIO_DIRECTORIO, values_only=True):
        if not fila or len(fila) <= COL["nombre"]:
            continue
        nombre = _texto(fila[COL["nombre"]])
        if not nombre:
            continue
        cotizo = (_texto(fila[COL["cotizo"]]) or "").upper() if len(fila) > COL["cotizo"] else ""
        filas.append({
            "grupo": _texto(fila[COL["grupo"]]),
            "nombre": nombre,
            "municipio": _texto(fila[COL["municipio"]]),
            "departamento": _texto(fila[COL["departamento"]]),
            "direccion": _texto(fila[COL["direccion"]]),
            "telefono": _texto(fila[COL["telefono"]]),
            "web_correo": _texto(fila[COL["web_correo"]]),
            "contacto": _texto(fila[COL["contacto"]]),
            "cotizo": 1 if cotizo.startswith("SI") else 0,
        })
    wb.close()
    return filas


def leer_bpr(ruta: Path) -> list[dict]:
    """Insumos con precio oficial del Banco de Precios de Referencia."""
    wb = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    hoja = next((h for h in HOJAS_BPR if h in wb.sheetnames), None)
    if not hoja:
        wb.close()
        raise ValueError(f"El archivo no tiene ninguna hoja de insumos {HOJAS_BPR}")

    filas = []
    for fila in wb[hoja].iter_rows(min_row=FILA_INICIO_BPR, values_only=True):
        if not fila or len(fila) <= COL_BPR["precio"]:
            continue
        codigo = _texto(fila[COL_BPR["codigo"]])
        nombre = _texto(fila[COL_BPR["nombre"]])
        precio = _numero(fila[COL_BPR["precio"]])
        if not codigo or not nombre or precio is None:
            continue
        filas.append({
            "codigo": codigo,
            "grupo": _texto(fila[COL_BPR["grupo"]]),
            "nombre": nombre,
            "unidad": _texto(fila[COL_BPR["unidad"]]),
            "precio": precio,
            "origen": _texto(fila[COL_BPR["origen"]]),
        })
    wb.close()
    return filas


def guardar_proveedores(filas: list[dict], periodo: Optional[str]) -> dict:
    """Inserta/actualiza proveedores y sus grupos. Devuelve un resumen."""
    proveedores, relaciones = {}, set()
    for f in filas:
        proveedores.setdefault(f["nombre"], f)  # primera aparición gana los datos de contacto
        if f["grupo"]:
            relaciones.add((f["nombre"], f["grupo"]))

    for nombre, datos in proveedores.items():
        execute_query(
            """INSERT INTO proveedor
                   (nombre, municipio, departamento, direccion, telefono, web_correo,
                    contacto, cotizo, periodo)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
               ON DUPLICATE KEY UPDATE
                   municipio = VALUES(municipio), departamento = VALUES(departamento),
                   direccion = VALUES(direccion), telefono = VALUES(telefono),
                   web_correo = VALUES(web_correo), contacto = VALUES(contacto),
                   cotizo = VALUES(cotizo), periodo = VALUES(periodo)""",
            (nombre, datos["municipio"], datos["departamento"], datos["direccion"],
             datos["telefono"], datos["web_correo"], datos["contacto"], datos["cotizo"], periodo),
            fetch=False,
        )

    ids = {r["nombre"]: r["id"] for r in (execute_query("SELECT id, nombre FROM proveedor") or [])}
    vinculos = 0
    for nombre, grupo in relaciones:
        pid = ids.get(nombre)
        if not pid:
            continue
        execute_query(
            "INSERT IGNORE INTO proveedor_grupo (proveedor_id, grupo) VALUES (%s, %s)",
            (pid, grupo), fetch=False,
        )
        vinculos += 1

    return {"proveedores": len(proveedores), "vinculos": vinculos}


def guardar_bpr(filas: list[dict], periodo: Optional[str]) -> dict:
    for f in filas:
        execute_query(
            """INSERT INTO insumo_referencia_idu
                   (codigo, grupo, nombre, unidad, precio, origen, periodo)
               VALUES (%s, %s, %s, %s, %s, %s, %s)
               ON DUPLICATE KEY UPDATE
                   grupo = VALUES(grupo), nombre = VALUES(nombre), unidad = VALUES(unidad),
                   precio = VALUES(precio), origen = VALUES(origen), periodo = VALUES(periodo)""",
            (f["codigo"], f["grupo"], f["nombre"], f["unidad"], f["precio"], f["origen"], periodo),
            fetch=False,
        )
    return {"insumos": len(filas)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--directorio", type=Path, help="Excel del directorio de proveedores")
    parser.add_argument("--bpr", type=Path, help="Excel del Banco de Precios de Referencia")
    parser.add_argument("--periodo", default=None, help="Etiqueta de la publicación, ej. '2026-I Fase I'")
    parser.add_argument("--dry-run", action="store_true", help="Solo informa, no escribe en la BD")
    args = parser.parse_args()

    if not args.directorio and not args.bpr:
        parser.error("Indica al menos --directorio o --bpr")

    if args.directorio:
        filas = leer_directorio(args.directorio)
        proveedores = {f["nombre"] for f in filas}
        grupos = {f["grupo"] for f in filas if f["grupo"]}
        log.info("Directorio: %d fila(s), %d proveedor(es), %d grupo(s)",
                 len(filas), len(proveedores), len(grupos))
        if not args.dry_run:
            log.info("Guardado: %s", guardar_proveedores(filas, args.periodo))

    if args.bpr:
        filas = leer_bpr(args.bpr)
        grupos = {f["grupo"] for f in filas if f["grupo"]}
        log.info("BPR: %d insumo(s) con precio, %d grupo(s)", len(filas), len(grupos))
        if not args.dry_run:
            log.info("Guardado: %s", guardar_bpr(filas, args.periodo))

    if args.dry_run:
        log.info("(DRY-RUN) No se escribió nada. Quita --dry-run para aplicar.")


if __name__ == "__main__":
    main()
